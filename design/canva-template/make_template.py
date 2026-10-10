"""
Sinh template tạp chí tiếng Nhật (khổ A4 dọc) cho nhóm nội dung chỉnh sửa trên Canva.

Mỗi trang được mô tả MỘT lần (danh sách khối: hình, chữ, chữ dọc…) rồi xuất ra:
  - .pptx  → nhập vào Canva: mỗi đoạn văn là một khung chữ, chữ dọc là khung chữ dọc thật
  - .pdf   → xem trước / đưa thẳng vào flipbook

Sách lật kiểu Nhật (phải → trái): trang chẵn nằm bên phải, trang lẻ bên trái (trừ bìa).
Không đánh số trang trong template: web tự in số trang.

Cài đặt (một lần):  pip install reportlab python-pptx
Font: Noto Sans JP, Noto Serif JP (bản đầy đủ, tĩnh 400/700), Playfair Display — đặt trong --fonts

Chạy:
  python make_template.py --fonts ./fonts --out ./non-template            # template 18 trang
  python make_template.py --fonts ./fonts --out ./sach-mau --demo         # bản demo ~100 trang
"""
import argparse
import json
import os

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt
from reportlab.lib.colors import HexColor
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

W, H = 595.28, 841.89  # A4, đơn vị point; toạ độ tính từ góc trên-trái

IVORY = '#f4ede0'
CREAM = '#ebe0cb'
INK = '#2a1f17'
INK2 = '#5e4d3d'
BROWN = '#3b2a1e'
BROWN2 = '#24180f'
GOLD = '#b8955a'
GOLD2 = '#d9bd84'
WINE = '#7d2f2a'
HAIR = '#d6c9b3'

# khoá font → (tên font trong Canva/PowerPoint, đậm?, file ttf cho PDF)
FONTS = {
    'sans': ('Noto Sans JP', False, 'NotoSansJP-Regular.ttf'),
    'sansB': ('Noto Sans JP', True, 'NotoSansJP-Bold.ttf'),
    'serif': ('Noto Serif JP', False, 'NotoSerifJP-Regular.ttf'),
    'serifB': ('Noto Serif JP', True, 'NotoSerifJP-Bold.ttf'),
    'display': ('Playfair Display', False, 'PlayfairDisplay-Regular.ttf'),
    'displayI': ('Playfair Display', False, 'PlayfairDisplay-Italic.ttf'),
}

# ───────────────────────────── nội dung mẫu
TITLE = 'ノン — 頭上の物語'
CHAPTERS = [
    ('起源', 'ノンの始まり'),
    ('村', '笠づくりの村'),
    ('素材', '葉と竹と糸'),
    ('技', '十六の輪'),
    ('暮らし', '日々の道具'),
    ('詩', '笠に描かれた詩'),
    ('旅', '市場から世界へ'),
    ('装い', 'アオザイとノン'),
    ('未来', '受け継ぐ手'),
    ('世界の帽子', '頭上の文化'),
]
ARTICLE_TITLES = ['かたちの理由', '手から手へ', '光と影', '雨の日のノン', '市場の朝', '次の世代へ']
PHOTO_TITLES = ['黄金の田んぼ', '夕暮れの道', '川辺の市場']
PARAS = [
    'ノンラー（ベトナムの葉の笠）は、強い日差しと突然のスコールから人々を守るために生まれました。円錐形のかたちは雨水をすばやく流し、広いつばは顔と肩に影をつくります。',
    '一つの笠を仕上げるには、熟練の職人でも丸一日かかります。竹を細く削って十六本の輪をつくり、乾燥させた椰子の葉を一枚ずつ重ねて、細い糸で縫いとめていきます。',
    'フエの「詩の笠」は、光にかざすと葉の間に詩や風景が浮かび上がります。薄い葉の層の間に切り絵を挟みこむ、繊細な技法です。',
    '田んぼで働く農家、市場の売り手、川を渡る船頭。ノンは日々の暮らしのなかで、日よけ、雨よけ、うちわ、そして水をすくう器にもなります。',
    '現代ではファッションや観光の象徴としても親しまれ、アオザイと組み合わせた姿はベトナムを代表するイメージの一つになりました。',
    '職人の数は年々減っていますが、若い世代が新しいデザインや素材に挑戦し、伝統を未来へとつないでいます。',
]
QUOTES = [
    '笠は頭の上にあるだけではない。\nかぶる人の物語を運んでいる。',
    '十六の輪に、\n村の時間が編みこまれている。',
    '光にかざせば、\n葉の奥から詩が現れる。',
]
WIKI = 'https://ja.wikipedia.org/wiki/%E3%83%8E%E3%83%B3%E3%83%A9%E3%83%BC'


# ───────────────────────────── khối dựng trang
def rect(x, y, w, h, fill, name=None):
    return {'k': 'rect', 'x': x, 'y': y, 'w': w, 'h': h, 'fill': fill, 'name': name}


def oval(x, y, w, h, fill, line=None, name=None):
    return {'k': 'oval', 'x': x, 'y': y, 'w': w, 'h': h, 'fill': fill, 'line': line, 'name': name}


def tri(x, y, w, h, fill, line=None, name=None, right=False):
    """tam giác cân, đỉnh hướng lên (right=True: hướng sang phải, như nút play)"""
    return {'k': 'tri', 'x': x, 'y': y, 'w': w, 'h': h, 'fill': fill, 'line': line, 'name': name, 'right': right}


def line(x1, y1, x2, y2, color, width=0.5):
    return {'k': 'line', 'x1': x1, 'y1': y1, 'x2': x2, 'y2': y2, 'color': color, 'width': width}


def text(x, y, w, h, s, font='sans', size=10, color=INK, align='left', spacing=1.0, vertical=False, link=None, tracking=0, name=None, para_gap=0):
    """para_gap: khoảng cách thêm giữa các đoạn (point)"""
    return {
        'k': 'text', 'x': x, 'y': y, 'w': w, 'h': h, 's': s, 'font': font, 'size': size, 'color': color,
        'align': align, 'spacing': spacing, 'vertical': vertical, 'link': link, 'tracking': tracking, 'name': name,
        'para_gap': para_gap,
    }


def group(name, *items):
    return {'k': 'group', 'name': name, 'items': list(items)}


def hat(cx, base_y, r, fill=GOLD, line_color=BROWN, rings=True, name='イラスト：ノン'):
    """ノンラー: 円錐 + 輪 + つば (cx = tâm, base_y = đáy nón)"""
    h = r * 0.78
    items = [tri(cx - r, base_y - h, 2 * r, h, fill, line_color)]
    if rings:
        for i in range(1, 8):
            t = i / 8
            yy = base_y - h * t
            items.append(line(cx - r * (1 - t), yy, cx + r * (1 - t), yy, line_color, 0.4))
    items.append(oval(cx - r, base_y - r * 0.08, 2 * r, r * 0.16, fill, line_color))
    return group(name, *items)


def sunset(x, y, w, h, seed, name='写真（差し替え）'):
    """Ảnh minh hoạ tạm (hoàng hôn trên ruộng) — nhóm nội dung thay bằng ảnh thật"""
    bands = ['#efc27a', '#d98a3d', '#9c4a2c', '#5a2d2a', '#2b1a22']
    # dải màu bầu trời: sáng ở chân trời, tối dần lên trên
    items = [rect(x, y + h * (0.58 - (i + 1) * 0.14), w, h * 0.15, col) for i, col in enumerate(bands)]
    items.append(rect(x, y, w, h * 0.06, bands[-1]))
    sx = x + w * (0.3 + 0.4 * ((seed * 37) % 10) / 10)
    items.append(oval(sx - h * 0.09, y + h * 0.41, h * 0.18, h * 0.18, '#f6dc9a'))
    items.append(rect(x, y + h * 0.58, w, h * 0.42, '#1b130e'))
    for i in range(14):
        items.append(line(x + w * 0.5, y + h * 0.58, x + w * (i / 13), y + h, '#3d2a1c', 0.6))
    items.append(hat(x + w * 0.62, y + h * 0.6, h * 0.11, fill='#120c08', line_color='#120c08', rings=False, name='シルエット'))
    return group(name, *items)


# ───────────────────────────── các kiểu trang
def chrome(dark=False):
    col = GOLD2 if dark else INK2
    return [
        text(40, 24, 300, 12, TITLE, 'sans', 7, col, tracking=1, name='ランニングヘッド'),
        line(40, 38, W - 40, 38, col, 0.4),
    ]


def page_cover(n, ctx):
    return BROWN2, [
        oval(W / 2 - W * 0.42, H * 0.48 - W * 0.42, W * 0.84, W * 0.84, BROWN),
        hat(W / 2, H * 0.57, W * 0.3, fill=GOLD, line_color=GOLD2),
        text(40, H * 0.17, W - 80, 100, 'NÓN', 'display', 84, IVORY, 'center', tracking=18, name='タイトル'),
        text(40, H * 0.30, W - 80, 36, TITLE, 'serifB', 24, GOLD2, 'center', name='サブタイトル'),
        text(40, H * 0.76, W - 80, 20, '日常の道具から、文化とファッションの象徴へ', 'sans', 10, IVORY, 'center', name='キャッチコピー'),
        text(40, H - 66, W - 80, 16, '第1号 — 2026年秋', 'sans', 9, GOLD, 'center', name='号数'),
    ]


def page_contents(n, ctx):
    items = chrome() + [text(W - 92, 64, 50, 160, '目次', 'serifB', 34, WINE, vertical=True, name='見出し（縦書き）')]
    rows = ctx['toc'][:10]
    y = 96
    for i, (title, kw, start) in enumerate(rows):
        items += [
            text(50, y, 44, 30, f'{i + 1:02d}', 'displayI', 22, GOLD),
            text(96, y + 2, 300, 22, title, 'serifB', 15, INK),
            text(96, y + 24, 300, 14, f'第{i + 1}章 · {kw}', 'sans', 8.5, INK2),
            text(W - 200, y + 6, 70, 16, f'{start:02d}', 'display', 10, INK2, 'right'),
            line(50, y + 46, W - 120, y + 46, HAIR, 0.5),
        ]
        y += 66
    return IVORY, items


def page_opener(n, ctx):
    ch = ctx['ch']
    kw, title = CHAPTERS[ch % len(CHAPTERS)]
    bg = BROWN if ch % 2 else WINE
    num_color = '#4a3628' if ch % 2 else '#8a3d37'  # số chương mờ: pha sẵn với màu nền
    # số chương nằm gọn trong trang, sát mép ngoài (trang chẵn bên phải, trang lẻ bên trái);
    # chừa góc dưới cho số trang web tự in
    right = n % 2 == 0
    num_x = W - 40 - 360 if right else 40
    return bg, [
        text(num_x, H - 330, 360, 280, f'{ch + 1:02d}', 'display', 240, num_color, 'right' if right else 'left', name='章番号（飾り）'),
        text(W - 130, 90, 70, H - 200, title, 'serifB', 44, IVORY, vertical=True, name='章タイトル（縦書き）'),
        text(50, 70, 200, 16, f'第{ch + 1}章', 'sansB', 11, GOLD2, name='章'),
        text(50, 90, 200, 14, kw, 'sans', 9, GOLD2, name='キーワード'),
    ] + chrome(dark=True)


def article_head(ch, k):
    return [
        text(50, 70, 300, 14, f'{ch + 1:02d} — {CHAPTERS[ch % len(CHAPTERS)][0]}', 'sansB', 9, GOLD, name='小見出し'),
        text(50, 88, W - 100, 40, ARTICLE_TITLES[(ch + k) % len(ARTICLE_TITLES)], 'serifB', 26, INK, name='見出し'),
    ]


def body(ch, k, count=3):
    return '\n'.join(PARAS[(ch + k + i) % len(PARAS)] for i in range(count))


def page_article_image(n, ctx):
    ch, k = ctx['ch'], ctx['k']
    return IVORY, chrome() + article_head(ch, k) + [
        group('画像（差し替え）', rect(50, 146, W - 100, 230, CREAM), hat(W / 2, 330, 110)),
        text(50, 384, W - 100, 14, '図：ノンラーの構造 — 十六本の竹の輪と椰子の葉', 'sans', 8, INK2, name='キャプション'),
        text(50, 420, W - 100, 300, body(ch, k), 'sans', 10.5, INK, spacing=1.55, name='本文', para_gap=8),
    ]


def page_article_text(n, ctx):
    ch, k = ctx['ch'], ctx['k']
    return IVORY, chrome() + article_head(ch, k) + [
        text(50, 150, W - 100, 440, body(ch, k, 5), 'sans', 10.5, INK, spacing=1.55, name='本文', para_gap=8),
        rect(50, 640, 3, 70, WINE, name='引用ライン'),
        text(66, 646, W - 130, 70, QUOTES[(ch + k) % len(QUOTES)].replace('\n', ''), 'serif', 14, WINE, spacing=1.4, name='引用'),
    ]


def page_vertical(n, ctx):
    ch = ctx['ch']
    return IVORY, chrome() + [
        text(W - 90, 66, 40, 340, CHAPTERS[ch % len(CHAPTERS)][1], 'serifB', 22, WINE, vertical=True, name='見出し（縦書き）'),
        text(50, 66, W - 160, H - 130, body(ch, 0, 5), 'serif', 11.5, INK, spacing=1.5, vertical=True, name='本文（縦書き）', para_gap=8),
    ]


def page_photo(n, ctx):
    ch = ctx['ch']
    return '#1b130e', [
        sunset(0, 0, W, H, ch),
        text(50, H - 190, W - 100, 44, PHOTO_TITLES[ch % len(PHOTO_TITLES)], 'serifB', 30, IVORY, name='写真タイトル'),
        text(50, H - 140, W - 100, 16, 'ベトナム中部', 'sans', 10, IVORY, name='写真キャプション'),
    ] + chrome(dark=True)


def page_quote(n, ctx):
    ch = ctx['ch']
    return CREAM, chrome() + [
        text(44, 120, 120, 150, '“', 'displayI', 140, GOLD, name='引用符'),
        text(70, H * 0.38, W - 140, 120, QUOTES[ch % len(QUOTES)], 'serifB', 22, BROWN, spacing=1.5, name='引用'),
        text(70, H * 0.38 + 130, W - 140, 16, '— フエの笠職人', 'sans', 9, INK2, name='出典'),
    ]


def page_video(n, ctx):
    # khung video: khớp overlays trong src/content/pdfBook.js (x 10%, y 30%, w 80%, h 45%)
    fx, fy, fw, fh = W * 0.10, H * 0.30, W * 0.80, H * 0.45
    return IVORY, chrome() + [
        text(50, 70, 300, 14, '映像', 'sansB', 9, GOLD, name='小見出し'),
        text(50, 88, W - 100, 40, '笠ができるまで', 'serifB', 26, INK, name='見出し'),
        text(50, 140, W - 100, 60, '竹を削り、葉を重ね、糸で縫いとめる。職人の一日を映像で追いました。', 'sans', 10.5, INK2, spacing=1.55, name='リード'),
        group('動画エリア', rect(fx, fy, fw, fh, BROWN2), tri(fx + fw / 2 - 18, fy + fh / 2 - 22, 44, 44, GOLD2, right=True)),
        text(50, fy + fh + 24, W - 100, 120, PARAS[1], 'sans', 10.5, INK, spacing=1.55, name='本文'),
    ]


def page_link(n, ctx):
    bg, items = page_article_text(n, {**ctx, 'k': 3})
    items = [it for it in items if it.get('name') not in ('引用ライン', '引用')]
    items.append(text(50, 660, W - 100, 20, 'もっと知る → ウィキペディア「ノンラー」', 'sansB', 12, WINE, link=WIKI, name='リンク'))
    return bg, items


def page_credits(n, ctx):
    roles = ['編集', 'デザイン', '翻訳', '写真', 'ウェブ']
    return BROWN2, chrome(dark=True) + [
        text(40, 190, W - 80, 30, '制作チーム', 'serifB', 22, GOLD2, 'center', name='見出し'),
        text(40, 250, W - 80, 200, '\n'.join(f'{r} — 氏名' for r in roles), 'sans', 11, IVORY, 'center', spacing=2.2, name='スタッフ'),
    ]


def page_back(n, ctx):
    return BROWN2, [
        hat(W / 2, H * 0.52, W * 0.12, fill=GOLD, line_color=GOLD2),
        text(40, H * 0.6, W - 80, 20, 'ありがとうございました', 'serif', 12, GOLD2, 'center', name='メッセージ'),
        text(40, H - 66, W - 80, 18, 'NÓN — 2026', 'display', 14, GOLD, 'center', tracking=4, name='ロゴ'),
    ]


def build_pages(demo):
    """Trả về danh sách (kiểu trang, ngữ cảnh). Luôn chẵn trang để bìa sau đứng riêng."""
    n_ch = 10 if demo else 2
    per = ['opener', 'image', 'vertical', 'photo', 'quote', 'text']
    if demo:
        per += ['image', 'text', 'quote']
    plan = [('cover', {}), ('contents', {})]
    toc = []
    for ch in range(n_ch):
        toc.append((CHAPTERS[ch][1], CHAPTERS[ch][0], len(plan) + 1))
        for k, kind in enumerate(per):
            plan.append((kind, {'ch': ch, 'k': k}))
        if ch == 0:
            plan += [('video', {'ch': 0}), ('link', {'ch': 0})]
    plan += [('credits', {}), ('back', {})]
    for _, ctx in plan:
        ctx['toc'] = toc
    return plan


KINDS = {
    'cover': page_cover, 'contents': page_contents, 'opener': page_opener, 'image': page_article_image,
    'text': page_article_text, 'vertical': page_vertical, 'photo': page_photo, 'quote': page_quote,
    'video': page_video, 'link': page_link, 'credits': page_credits, 'back': page_back,
}


# ───────────────────────────── xuất PDF
def wrap(s, font, size, width, tracking):
    """Trả về các dòng; None = ranh giới giữa hai đoạn"""
    lines = []
    for i, para in enumerate(s.split('\n')):
        if i:
            lines.append(None)
        cur = ''
        for ch in para:
            if pdfmetrics.stringWidth(cur + ch, font, size) + tracking * len(cur) > width and cur:
                if ch in '、。」）':  # 禁則: không để dấu câu đứng đầu dòng
                    cur += ch
                    continue
                lines.append(cur)
                cur = ch
            else:
                cur += ch
        lines.append(cur)
    return lines


def pdf_draw(c, it):
    k = it['k']
    if k == 'group':
        for sub in it['items']:
            pdf_draw(c, sub)
        return
    if k == 'line':
        c.setStrokeColor(HexColor(it['color']))
        c.setLineWidth(it['width'])
        c.line(it['x1'], H - it['y1'], it['x2'], H - it['y2'])
        return
    if k in ('rect', 'oval', 'tri'):
        c.setFillColor(HexColor(it['fill']))
        stroke = 1 if it.get('line') else 0
        if stroke:
            c.setStrokeColor(HexColor(it['line']))
            c.setLineWidth(1)
        x, y, w, h = it['x'], H - it['y'] - it['h'], it['w'], it['h']
        if k == 'rect':
            c.rect(x, y, w, h, stroke=0, fill=1)
        elif k == 'oval':
            c.ellipse(x, y, x + w, y + h, stroke=stroke, fill=1)
        else:
            p = c.beginPath()
            if it.get('right'):
                p.moveTo(x, y)
                p.lineTo(x, y + h)
                p.lineTo(x + w, y + h / 2)
            else:
                p.moveTo(x, y)
                p.lineTo(x + w / 2, y + h)
                p.lineTo(x + w, y)
            p.close()
            c.drawPath(p, stroke=stroke, fill=1)
        return
    # chữ
    font, size = it['font'], it['size']
    c.setFillColor(HexColor(it['color']))
    lead = size * 1.2 * it['spacing']
    if it['vertical']:
        pdf_vertical(c, it, font, size, lead)
        return
    y = H - it['y'] - size * 0.92
    for ln in wrap(it['s'], font, size, it['w'], it['tracking']):
        if ln is None:
            y -= it['para_gap']
            continue
        width = pdfmetrics.stringWidth(ln, font, size) + it['tracking'] * max(len(ln) - 1, 0)
        x = it['x'] + {'left': 0, 'center': (it['w'] - width) / 2, 'right': it['w'] - width}[it['align']]
        t = c.beginText(x, y)
        t.setFont(font, size)
        t.setCharSpace(it['tracking'])
        t.textOut(ln)
        c.drawText(t)
        if it['link']:
            c.linkURL(it['link'], (x - 2, y - 4, x + width + 2, y + size), relative=0)
            c.setStrokeColor(HexColor(it['color']))
            c.setLineWidth(0.6)
            c.line(x, y - 3, x + width, y - 3)
        y -= lead


def pdf_vertical(c, it, font, size, lead):
    """縦書き: trên → dưới, cột từ phải sang trái"""
    c.setFont(font, size)
    top = H - it['y']
    x = it['x'] + it['w'] - size
    for para in it['s'].split('\n'):
        y = top - size
        for ch in para:
            if y < top - it['h']:
                x -= lead
                y = top - size
            if ch in '、。':
                c.drawString(x + size * 0.6, y + size * 0.6, ch)
            elif ch in 'ー—「」（）':
                c.saveState()
                c.translate(x + size / 2, y + size * 0.38)
                c.rotate(-90)
                c.drawCentredString(0, -size * 0.36, ch)
                c.restoreState()
            else:
                c.drawString(x, y, ch)
            y -= size
        x -= lead + it['para_gap']


def write_pdf(path, plan, fonts_dir):
    for key, (_, _, file) in FONTS.items():
        pdfmetrics.registerFont(TTFont(key, os.path.join(fonts_dir, file)))
    c = canvas.Canvas(path, pagesize=(W, H))
    c.setTitle(TITLE)
    for n, (kind, ctx) in enumerate(plan, start=1):
        bg, items = KINDS[kind](n, ctx)
        c.setFillColor(HexColor(bg))
        c.rect(0, 0, W, H, stroke=0, fill=1)
        for it in items:
            pdf_draw(c, it)
        c.showPage()
    c.save()


# ───────────────────────────── xuất PPTX (để nhập vào Canva)
def rgb(hex_):
    return RGBColor.from_string(hex_.lstrip('#'))


def pptx_shape(shapes, it):
    k = it['k']
    if k == 'group':
        g = shapes.add_group_shape()
        for sub in it['items']:
            pptx_shape(g.shapes, sub)
        g.name = it['name']
        return
    if k == 'line':
        ln = shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Pt(it['x1']), Pt(it['y1']), Pt(it['x2']), Pt(it['y2']))
        ln.line.color.rgb = rgb(it['color'])
        ln.line.width = Pt(it['width'])
        return
    if k in ('rect', 'oval', 'tri'):
        kind = {'rect': MSO_SHAPE.RECTANGLE, 'oval': MSO_SHAPE.OVAL, 'tri': MSO_SHAPE.ISOSCELES_TRIANGLE}[k]
        sh = shapes.add_shape(kind, Pt(it['x']), Pt(it['y']), Pt(it['w']), Pt(it['h']))
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(it['fill'])
        if it.get('line'):
            sh.line.color.rgb = rgb(it['line'])
            sh.line.width = Pt(1)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        if it.get('right'):
            sh.rotation = 90
        if it.get('name'):
            sh.name = it['name']
        return
    # chữ
    family, bold, _ = FONTS[it['font']]
    tb = shapes.add_textbox(Pt(it['x']), Pt(it['y']), Pt(it['w']), Pt(it['h']))
    if it.get('name'):
        tb.name = it['name']
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for side in ('left', 'right', 'top', 'bottom'):
        setattr(tf, f'margin_{side}', 0)
    if it['vertical']:
        tf._txBody.find(qn('a:bodyPr')).set('vert', 'eaVert')
    for i, para in enumerate(it['s'].split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT}[it['align']]
        p.line_spacing = it['spacing']
        if it['para_gap']:
            p.space_after = Pt(it['para_gap'])
        r = p.add_run()
        r.text = para
        f = r.font
        f.size = Pt(it['size'])
        f.bold = bold
        f.italic = it['font'] == 'displayI'
        f.color.rgb = rgb(it['color'])
        f.name = family
        rpr = r._r.get_or_add_rPr()
        for tag in ('a:ea', 'a:cs'):  # font cho chữ Nhật (East Asian)
            el = rpr.find(qn(tag))
            if el is None:
                el = rpr.makeelement(qn(tag), {})
                rpr.append(el)
            el.set('typeface', family)
        if it['tracking']:
            rpr.set('spc', str(int(it['tracking'] * 100)))
        if it['link']:
            r.hyperlink.address = it['link']


def write_pptx(path, plan):
    prs = Presentation()
    prs.slide_width = Pt(W)
    prs.slide_height = Pt(H)
    blank = prs.slide_layouts[6]
    for n, (kind, ctx) in enumerate(plan, start=1):
        bg, items = KINDS[kind](n, ctx)
        slide = prs.slides.add_slide(blank)
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = rgb(bg)
        for it in items:
            pptx_shape(slide.shapes, it)
    prs.save(path)


# ───────────────────────────── xuất JSON (dữ liệu trang cho editor trên web)
FONT_JSON = {
    'sans': ('Noto Sans JP', 400, False), 'sansB': ('Noto Sans JP', 700, False),
    'serif': ('Noto Serif JP', 400, False), 'serifB': ('Noto Serif JP', 700, False),
    'display': ('Playfair Display', 400, False), 'displayI': ('Playfair Display', 400, True),
}
KIND_NAMES = {
    'cover': 'Bìa trước', 'contents': 'Mục lục', 'opener': 'Mở chương', 'image': 'Bài có ảnh',
    'vertical': 'Bài chữ dọc', 'photo': 'Ảnh tràn trang', 'quote': 'Trích dẫn', 'text': 'Bài chỉ có chữ',
    'video': 'Trang video', 'link': 'Bài có link', 'credits': 'Ê-kíp', 'back': 'Bìa sau',
}


def bbox(it):
    k = it['k']
    if k == 'group':
        boxes = [bbox(c) for c in it['items']]
        x1 = min(b[0] for b in boxes)
        y1 = min(b[1] for b in boxes)
        return x1, y1, max(b[0] + b[2] for b in boxes) - x1, max(b[1] + b[3] for b in boxes) - y1
    if k == 'line':
        x, y = min(it['x1'], it['x2']), min(it['y1'], it['y2'])
        return x, y, abs(it['x2'] - it['x1']), abs(it['y2'] - it['y1'])
    return it['x'], it['y'], it['w'], it['h']


def r2(v):
    return round(v, 2)


def to_json(it, ox=0, oy=0):
    """Khối của template → phần tử JSON; (ox, oy) = gốc toạ độ của nhóm cha"""
    k = it['k']
    if k == 'group' and it['name'] == '動画エリア':
        # khung video của template → phần tử video thật trên web (gắn sẵn video demo)
        x, y, w, h = bbox(it)
        return {'type': 'video', 'name': 'Video', 'x': r2(x - ox), 'y': r2(y - oy), 'w': r2(w), 'h': r2(h),
                'src': 'assets/videos/making-hat.mp4'}
    if k == 'group':
        x, y, w, h = bbox(it)
        return {'type': 'group', 'name': it['name'], 'x': r2(x - ox), 'y': r2(y - oy), 'w': r2(w), 'h': r2(h),
                'bw': r2(w), 'bh': r2(h), 'children': [to_json(c, x, y) for c in it['items']]}
    if k == 'line':
        x, y, w, h = bbox(it)
        # đường kẻ ngang/dọc: cho hộp dày tối thiểu để chọn được
        pad_y = 2 if h < 2 else 0
        pad_x = 2 if w < 2 else 0
        bx, by, bw, bh = x - pad_x, y - pad_y, w + 2 * pad_x, h + 2 * pad_y
        return {'type': 'line', 'x': r2(bx - ox), 'y': r2(by - oy), 'w': r2(bw), 'h': r2(bh),
                'x1': r2((it['x1'] - bx) / bw), 'y1': r2((it['y1'] - by) / bh),
                'x2': r2((it['x2'] - bx) / bw), 'y2': r2((it['y2'] - by) / bh),
                'stroke': it['color'], 'strokeWidth': it['width']}
    if k in ('rect', 'oval', 'tri'):
        el = {'type': 'shape', 'shape': {'rect': 'rect', 'oval': 'ellipse', 'tri': 'triangle'}[k],
              'x': r2(it['x'] - ox), 'y': r2(it['y'] - oy), 'w': r2(it['w']), 'h': r2(it['h']), 'fill': it['fill']}
        if it.get('line'):
            el['stroke'] = it['line']
            el['strokeWidth'] = 1
        if it.get('right'):
            el['rot'] = 90
        if it.get('name'):
            el['name'] = it['name']
        return el
    family, weight, italic = FONT_JSON[it['font']]
    el = {'type': 'text', 'x': r2(it['x'] - ox), 'y': r2(it['y'] - oy), 'w': r2(it['w']), 'h': r2(it['h']),
          'text': it['s'], 'font': family, 'weight': weight, 'italic': italic, 'size': it['size'],
          'color': it['color'], 'align': it['align'], 'lineHeight': r2(1.2 * it['spacing'])}
    for key, out in (('vertical', 'vertical'), ('link', 'link'), ('tracking', 'tracking'), ('para_gap', 'paraGap'), ('name', 'name')):
        if it.get(key):
            el[out] = it[key]
    return el


def write_json(path, plan):
    pages = []
    for n, (kind, ctx) in enumerate(plan, start=1):
        bg, items = KINDS[kind](n, ctx)
        pages.append({'template': kind, 'templateName': KIND_NAMES[kind], 'bg': bg, 'elements': [to_json(it) for it in items]})
    with open(path, 'w', encoding='utf-8') as f:
        json.dump({'width': W, 'height': H, 'pages': pages}, f, ensure_ascii=False, indent=1)


if __name__ == '__main__':
    ap = argparse.ArgumentParser()
    ap.add_argument('--fonts', required=True)
    ap.add_argument('--out', required=True, help='đường dẫn không đuôi; sinh ra .pdf và .pptx')
    ap.add_argument('--demo', action='store_true', help='bản dài ~100 trang để thử flipbook')
    ap.add_argument('--no-pptx', action='store_true')
    ap.add_argument('--json', help='xuất thêm dữ liệu trang cho editor web (.json)')
    a = ap.parse_args()
    plan = build_pages(a.demo)
    if a.json:
        write_json(a.json, plan)
    write_pdf(a.out + '.pdf', plan, a.fonts)
    if not a.no_pptx:
        write_pptx(a.out + '.pptx', plan)
    print(f'{len(plan)} trang →', a.out)
    for n, (kind, ctx) in enumerate(plan, start=1):
        if kind in ('opener', 'video', 'link'):
            print(f'  trang {n}: {kind}' + (f' (chương {ctx["ch"] + 1})' if kind == 'opener' else ''))
